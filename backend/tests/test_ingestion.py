import unittest
import sys
import tempfile
import os
from pathlib import Path

backend_dir = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(backend_dir))

from app.services.ingestion.extractor import extract_text, UnsupportedFormatError
import docx
import fitz

class TestIngestion(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.gettempdir()

    def test_plain_text_extraction(self):
        txt_path = os.path.join(self.temp_dir, "test_doc.txt")
        sample_content = "Clause 1: The tenant shall pay $1,500 rent monthly on the 1st day."
        with open(txt_path, "w", encoding="utf-8") as f:
            f.write(sample_content)

        text, page_count = extract_text(txt_path, "txt")
        self.assertIn("Clause 1:", text)
        self.assertGreaterEqual(page_count, 1)
        os.remove(txt_path)

    def test_docx_extraction(self):
        docx_path = os.path.join(self.temp_dir, "test_doc.docx")
        doc = docx.Document()
        doc.add_paragraph("Section 1: Security Deposit")
        doc.add_paragraph("Tenant shall deposit $2,000 upon signing.")
        
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Utility"
        table.cell(0, 1).text = "Responsible Party"
        table.cell(1, 0).text = "Water"
        table.cell(1, 1).text = "Landlord"
        
        doc.save(docx_path)

        text, page_count = extract_text(docx_path, "docx")
        self.assertIn("Security Deposit", text)
        self.assertIn("Water | Landlord", text)
        os.remove(docx_path)

    def test_digital_pdf_extraction(self):
        pdf_path = os.path.join(self.temp_dir, "test_doc.pdf")
        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((50, 50), "ARTICLE 1: TERM AND TERMINATION\nThis lease terminates in 30 days.")
        doc.save(pdf_path)
        doc.close()

        text, page_count = extract_text(pdf_path, "pdf")
        self.assertIn("TERM AND TERMINATION", text)
        self.assertEqual(page_count, 1)
        os.remove(pdf_path)

    def test_unsupported_format(self):
        fake_path = os.path.join(self.temp_dir, "test.xyz")
        with self.assertRaises(UnsupportedFormatError):
            extract_text(fake_path, "xyz")

if __name__ == "__main__":
    unittest.main()
